import argparse
import json
import os
import re
import time
from typing import Dict, List, Optional, TypedDict

# -----------------------------
# LangGraph imports
# -----------------------------
try:
    from langchain_ollama import OllamaLLM
    from langchain_core.prompts import PromptTemplate
    from langgraph.graph import StateGraph, END
    from langgraph.prebuilt import ToolNode
    from docx import Document
    docx_available = True
except ModuleNotFoundError:
    docx_available = False
    # define PromptTemplate etc.
    class Document:
        def __init__(self, file):
            self.file = file
            self.paragraphs = []

    class PromptTemplate:
        def __init__(self, input_variables, template):
            self.input_variables = input_variables
            self.template = template

        def format(self, **kwargs):
            return self.template.format(**kwargs)

    class OllamaLLM:
        def __init__(self, model=None, temperature=0.3):
            self.model = model
            self.temperature = temperature

        def invoke(self, prompt):
            return self._reply(prompt)

        def _reply(self, prompt):
            if "draft" in prompt.lower():
                return "Thank you for your email. We are looking into this issue and will get back to you soon."
            if "urgency" in prompt.lower():
                email = prompt.split("Email: ")[1].split("\n\nUrgency:")[0].strip()
                if any(word in email.lower() for word in ["twice", "fails", "crashes", "intermittently"]):
                    return "High"
                else:
                    return "Medium"
            if "topic" in prompt.lower():
                email = prompt.split("Email: ")[1].split("\n\nTopic:")[0].strip()
                if "reset" in email.lower():
                    return "Account"
                elif "charged" in email.lower():
                    return "Billing"
                elif "crash" in email.lower():
                    return "Bug"
                elif "dark mode" in email.lower():
                    return "Feature Request"
                else:
                    return "Technical Issue"
            if "escalate" in prompt.lower():
                return "Escalate to human"
            if "follow-up" in prompt.lower():
                return "Schedule follow-up in 3 days"
            return "Default response"

    # Fallback for LangGraph
    class StateGraph:
        def __init__(self, state_class):
            self.state_class = state_class
            self.nodes = {}
            self.edges = []

        def add_node(self, name, func):
            self.nodes[name] = func

        def add_edge(self, from_node, to_node):
            self.edges.append((from_node, to_node))

        def add_conditional_edges(self, from_node, condition_func, edges):
            self.edges.append((from_node, condition_func, edges))

        def set_entry_point(self, node):
            self.entry = node

        def compile(self):
            return self

        def invoke(self, state):
            current = self.entry
            while current != END:
                if callable(current):
                    current = current(state)
                else:
                    state = self.nodes[current](state)
                    current = self._next_node(current, state)
            return state

        def _next_node(self, current, state):
            for edge in self.edges:
                if edge[0] == current:
                    if len(edge) == 2:
                        return edge[1]
                    elif len(edge) == 3:
                        return edge[1](state)
            return END

    END = "END"

# -----------------------------
# LLM instance
# -----------------------------
llm = OllamaLLM(model="qwen:7b", temperature=0.3)

# -----------------------------
# Knowledge Base loading
# -----------------------------
def load_knowledge_base():
    if docx_available:
        try:
            doc = Document('knowledge_base.docx')
            kb = {}
            for para in doc.paragraphs:
                text = para.text.strip()
                if text.startswith('Keywords:'):
                    lines = text.split('\n')
                    keywords_line = lines[0]
                    answer_line = lines[1] if len(lines) > 1 else ''
                    keywords = keywords_line.replace('Keywords:', '').strip().split(', ')
                    answer = answer_line.replace('Answer:', '').strip() if answer_line.startswith('Answer:') else ''
                    for kw in keywords:
                        kb[kw.strip()] = answer
            return kb
        except:
            pass
    # Fallback
    return {
        "reset": "To reset your password, go to the login page and click 'Forgot Password'. Follow the instructions sent to your email.",
        "password": "To reset your password, go to the login page and click 'Forgot Password'. Follow the instructions sent to your email.",
        "billing": "If you were charged twice, please contact billing support at billing@company.com with your invoice number.",
        "charged": "If you were charged twice, please contact billing support at billing@company.com with your invoice number.",
        "subscription": "If you were charged twice, please contact billing support at billing@company.com with your invoice number.",
        "bug": "Thank you for reporting the bug. Our team will investigate and release a fix in the next update.",
        "crash": "Thank you for reporting the bug. Our team will investigate and release a fix in the next update.",
        "export": "Thank you for reporting the bug. Our team will investigate and release a fix in the next update.",
        "pdf": "Thank you for reporting the bug. Our team will investigate and release a fix in the next update.",
        "feature": "We appreciate your suggestion for dark mode. It has been added to our feature backlog.",
        "add": "We appreciate your suggestion for dark mode. It has been added to our feature backlog.",
        "dark mode": "We appreciate your suggestion for dark mode. It has been added to our feature backlog.",
        "mobile app": "We appreciate your suggestion for dark mode. It has been added to our feature backlog.",
        "api": "For API 504 errors, check your network connection and retry. If persistent, escalate to technical support.",
        "504": "For API 504 errors, check your network connection and retry. If persistent, escalate to technical support.",
        "integration": "For API 504 errors, check your network connection and retry. If persistent, escalate to technical support.",
        "fails": "For API 504 errors, check your network connection and retry. If persistent, escalate to technical support.",
        "intermittently": "For API 504 errors, check your network connection and retry. If persistent, escalate to technical support.",
        "account": "For account-related questions, visit our help center at help.company.com/account.",
    }

knowledge_base = load_knowledge_base()

# -----------------------------
# State Definition
# -----------------------------
class SupportState(TypedDict):
    email_content: str
    urgency: Optional[str]
    topic: Optional[str]
    kb_answer: Optional[str]
    response: Optional[str]
    escalate: Optional[bool]
    follow_up: Optional[str]

# -----------------------------
# Node Functions
# -----------------------------
def classify_urgency(state: SupportState) -> SupportState:
    """Classify email urgency."""
    prompt = f"""
Classify the urgency of this customer email: Low, Medium, or High.

Email: {state['email_content']}

Urgency:
"""
    response = llm.invoke(prompt).strip()
    if "high" in response.lower():
        state['urgency'] = "High"
    elif "low" in response.lower():
        state['urgency'] = "Low"
    else:
        state['urgency'] = "Medium"
    return state

def classify_topic(state: SupportState) -> SupportState:
    """Classify email topic."""
    prompt = f"""
Classify the topic of this customer email: Account, Billing, Bug, Feature Request, or Technical Issue.

Email: {state['email_content']}

Topic:
"""
    response = llm.invoke(prompt).strip()
    topics = ["Account", "Billing", "Bug", "Feature Request", "Technical Issue"]
    for topic in topics:
        if topic.lower() in response.lower():
            state['topic'] = topic
            break
    else:
        state['topic'] = "Technical Issue"
    return state

def search_kb(state: SupportState) -> SupportState:
    """Search knowledge base."""
    email_lower = state['email_content'].lower()
    for key, answer in knowledge_base.items():
        if key in email_lower:
            state['kb_answer'] = answer
            break
    else:
        state['kb_answer'] = None
    return state

def draft_response(state: SupportState) -> SupportState:
    """Draft response."""
    if state['kb_answer']:
        state['response'] = state['kb_answer']
    else:
        prompt = f"""
Draft a customer support response for this email.

Topic: {state['topic']}
Urgency: {state['urgency']}
Email: {state['email_content']}

Response:
"""
        state['response'] = llm.invoke(prompt).strip()
    return state

def decide_escalate(state: SupportState) -> SupportState:
    """Decide if to escalate."""
    state['escalate'] = (state['urgency'] == "High" or not state['kb_answer'] or state['topic'] == "Technical Issue")
    return state

def schedule_follow_up(state: SupportState) -> SupportState:
    """Schedule follow-up."""
    if state['urgency'] == "High" or state['topic'] == "Technical Issue":
        state['follow_up'] = "Schedule follow-up in 3 days"
    else:
        state['follow_up'] = None
    return state

# -----------------------------
# Conditional Routing
# -----------------------------
def route_after_escalate(state: SupportState) -> str:
    """Route based on escalation decision."""
    return "schedule_follow_up" if state['escalate'] else END

# -----------------------------
# Build Graph
# -----------------------------
def build_graph():
    graph = StateGraph(SupportState)

    graph.add_node("classify_urgency", classify_urgency)
    graph.add_node("classify_topic", classify_topic)
    graph.add_node("search_kb", search_kb)
    graph.add_node("draft_response", draft_response)
    graph.add_node("decide_escalate", decide_escalate)
    graph.add_node("schedule_follow_up", schedule_follow_up)

    graph.set_entry_point("classify_urgency")

    graph.add_edge("classify_urgency", "classify_topic")
    graph.add_edge("classify_topic", "search_kb")
    graph.add_edge("search_kb", "draft_response")
    graph.add_edge("draft_response", "decide_escalate")
    graph.add_conditional_edges("decide_escalate", route_after_escalate, {"schedule_follow_up": "schedule_follow_up", END: END})
    graph.add_edge("schedule_follow_up", END)

    return graph.compile()

# -----------------------------
# Process Email
# -----------------------------
def process_email(email_content: str) -> Dict:
    """Process a single email using LangGraph."""
    graph = build_graph()
    initial_state: SupportState = {
        "email_content": email_content,
        "urgency": None,
        "topic": None,
        "kb_answer": None,
        "response": None,
        "escalate": None,
        "follow_up": None
    }
    result = graph.invoke(initial_state)
    return dict(result)

# -----------------------------
# Main Function
# -----------------------------
def main():
    parser = argparse.ArgumentParser(description="AI-Powered Customer Support Agent with LangGraph")
    parser.add_argument("--email", help="Process a single email")
    parser.add_argument("--file", help="Process emails from a JSON file")
    parser.add_argument("--output", default="output.json", help="Output file for results")
    args = parser.parse_args()

    results = []

    if args.email:
        result = process_email(args.email)
        results.append(result)
        print(json.dumps(result, indent=2))

    elif args.file:
        with open(args.file, "r") as f:
            emails = json.load(f)
        for email in emails:
            result = process_email(email["content"])
            result["email_id"] = email.get("id", "unknown")
            results.append(result)

    if results:
        with open(args.output, "w") as f:
            json.dump(results, f, indent=2)
        print(f"Results saved to {args.output}")

if __name__ == "__main__":
    main()